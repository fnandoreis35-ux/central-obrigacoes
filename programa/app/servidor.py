"""Servidor local: entrega a tela e responde às ações (banco, guias, e-mail, importação)."""
import base64
import csv
import io
import json
import mimetypes
import os
import re
import threading
import traceback
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import unquote, urlparse

from . import email_srv, guias
from .banco import Banco, agora_iso

AQUI = os.path.dirname(os.path.abspath(__file__))
WEB = os.path.join(AQUI, "web")
COLECOES = {"clientes", "obrigacoes", "processos", "tarefas", "config"}
REGIMES = ["Simples Nacional", "MEI", "Lucro Presumido", "Lucro Real", "Imune/Isenta", "Pessoa Física"]


class Erro(Exception):
    def __init__(self, msg, status=400):
        super().__init__(msg)
        self.status = status


class App:
    def __init__(self, banco: Banco, pasta_dados, pasta_usuario, versao, salvar_pasta_dados=None):
        self.banco = banco
        self.pasta_dados = pasta_dados
        self.pasta_usuario = pasta_usuario
        self.versao = versao
        self.cfg_email = email_srv.ConfigEmail(pasta_usuario)
        self.salvar_pasta_dados = salvar_pasta_dados
        self.lock_envio = threading.Lock()

    # ---------------- utilidades ----------------
    def _email_cfg(self):
        c = self.cfg_email.ler(com_senha=True)
        if not c.get("remetente") or not c.get("smtp") or not c.get("senha"):
            raise Erro("Configure o e-mail de envio em Configurações > E-mail antes de continuar.")
        return c

    def _pasta_guias(self, comp):
        if not re.fullmatch(r"\d{4}-\d{2}", comp or ""):
            raise Erro("Competência inválida.")
        p = os.path.join(self.pasta_dados, "guias", comp)
        os.makedirs(p, exist_ok=True)
        return p

    # ---------------- guias ----------------
    def analisar_guias(self, corpo):
        comp = corpo.get("comp")
        pasta = self._pasta_guias(comp)
        clientes = self.banco.listar("clientes")
        obrig = {o["id"]: o for o in self.banco.listar("obrigacoes")}
        tarefas = {t["id"]: t for t in self.banco.listar("tarefas")}
        saida = []
        for a in corpo.get("arquivos") or []:
            nome = os.path.basename(a["nome"])
            if not nome.lower().endswith(".pdf"):
                continue
            caminho = os.path.join(pasta, nome)
            with open(caminho, "wb") as f:
                f.write(base64.b64decode(a["b64"]))
            linha = {"arquivo": nome, "alertas": []}
            try:
                info = guias.extrair(guias.texto_pdf(caminho))
            except Exception as e:  # noqa: BLE001
                info = {"cnpj": None, "tipo": None, "valor": None, "vencimento": None, "periodo": None}
                linha["alertas"].append(f"Não consegui ler o PDF ({e}).")
            linha.update(info)
            c = guias.casar(info, clientes)
            if not c:
                linha["alertas"].append("Cliente não encontrado pelo CNPJ.")
            else:
                linha["clienteId"], linha["clienteNome"] = c["id"], c["nome"]
                if not c.get("emails"):
                    linha["alertas"].append("Cliente sem e-mail de entrega.")
            if not info.get("tipo"):
                linha["alertas"].append("Tipo da guia não identificado. Escolha a obrigação.")
            elif info["tipo"] not in obrig:
                linha["alertas"].append(f"Obrigação '{info['tipo']}' não existe no catálogo.")
            else:
                linha["obrigId"] = info["tipo"]
                if not obrig[info["tipo"]].get("enviaCliente"):
                    linha["alertas"].append("Obrigação marcada para não enviar ao cliente.")
            if c and linha.get("obrigId"):
                t = tarefas.get(f"{comp}_{c['id']}")
                it = (t or {}).get("itens", {}).get(linha["obrigId"])
                if not it:
                    linha["alertas"].append(f"Não há tarefa desta obrigação em {comp[5:]}/{comp[:4]}.")
                else:
                    linha["vencPrevisto"] = it.get("venc")
                    if info.get("vencimento") and it.get("venc") and info["vencimento"] != it["venc"]:
                        linha["alertas"].append(
                            f"Vencimento da guia ({guias.br(info['vencimento'])}) diferente do previsto ({guias.br(it['venc'])}).")
                    if it.get("status") in ("enviada", "ciente"):
                        linha["alertas"].append("Esta obrigação já foi enviada ao cliente nesta competência.")
            saida.append(linha)
        return {"guias": saida}

    def _corpo_email(self, cfg, cliente, comp, protocolo, itens_txt):
        escritorio = cfg.get("nome") or ""
        compbr = comp[5:] + "/" + comp[:4]
        subs = {"{cliente}": cliente["nome"], "{cnpj}": cliente.get("cnpj") or "", "{competencia}": compbr,
                "{protocolo}": protocolo, "{escritorio}": escritorio}

        def aplica(s):
            for k, v in subs.items():
                s = s.replace(k, v)
            return s

        assunto = aplica(cfg.get("modeloAssunto") or "Guias {competencia} | {cliente} | Protocolo {protocolo}")
        if protocolo not in assunto:
            assunto += f" | Protocolo {protocolo}"
        corpo = aplica(cfg.get("modeloCorpo") or "Olá, {cliente}.\n\nSeguem as guias de {competencia}:\n\n{guias}")
        assin = aplica(cfg.get("assinatura") or "")
        lista_txt = "\n".join("- " + x for x in itens_txt)
        texto = corpo.replace("{guias}", lista_txt) + ("\n\n" + assin if assin else "")

        def esc(s):
            return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        partes = []
        for bloco in corpo.split("\n\n"):
            if "{guias}" in bloco:
                antes, _, depois = bloco.partition("{guias}")
                if antes.strip():
                    partes.append("<p>" + esc(antes.strip()).replace("\n", "<br>") + "</p>")
                partes.append("<ul>" + "".join(f"<li>{esc(x)}</li>" for x in itens_txt) + "</ul>")
                if depois.strip():
                    partes.append("<p>" + esc(depois.strip()).replace("\n", "<br>") + "</p>")
            elif bloco.strip():
                partes.append("<p>" + esc(bloco.strip()).replace("\n", "<br>") + "</p>")
        if assin:
            partes.append("<p>" + esc(assin).replace("\n", "<br>") + "</p>")
        return assunto, "\n".join(partes), texto

    def enviar_guias(self, corpo):
        comp = corpo.get("comp")
        pasta = self._pasta_guias(comp)
        ec = self._email_cfg()
        cfg = self.banco.obter("config", "escritorio") or {}
        obrig = {o["id"]: o for o in self.banco.listar("obrigacoes")}
        resultados = []
        with self.lock_envio:
            for g in corpo.get("grupos") or []:
                cli = self.banco.obter("clientes", g["clienteId"])
                if not cli:
                    resultados.append({"clienteId": g["clienteId"], "ok": False, "erro": "Cliente não encontrado."})
                    continue
                cli["id"] = g["clienteId"]
                para = cli.get("emails") or []
                if not para:
                    resultados.append({"cliente": cli["nome"], "ok": False, "erro": "Cliente sem e-mail."})
                    continue
                tid = f"{comp}_{cli['id']}"
                t = self.banco.obter("tarefas", tid)
                if not t:
                    t = {"competencia": comp, "clienteId": cli["id"], "clienteNome": cli["nome"], "criadoEm": agora_iso(),
                         "itens": {}, "envios": []}
                    self.banco.set("tarefas", tid, t)
                todas = self.banco.listar("tarefas")
                seq, pref = guias.proximo_protocolo(comp, todas)
                protocolo = f"{pref}{seq:04d}"
                itens_txt, anexos, nomes = [], [], []
                for it in g["itens"]:
                    ob = obrig.get(it["obrigId"], {"sigla": it["obrigId"], "nome": ""})
                    venc = it.get("vencimento") or (t["itens"].get(it["obrigId"]) or {}).get("venc")
                    itens_txt.append(f"{ob['sigla']} – {ob['nome']} – vence {guias.br(venc)} – {guias.moeda(it.get('valor'))}")
                    caminho = os.path.join(pasta, os.path.basename(it["arquivo"]))
                    with open(caminho, "rb") as f:
                        dados = f.read()
                    nome_anexo = re.sub(r"[^A-Za-z0-9_.-]+", "", f"{ob['sigla']}_{comp[5:]}-{comp[:4]}_{cli['nome'][:30]}".replace(" ", "_").replace("/", "-")) + ".pdf"
                    while nome_anexo in nomes:
                        nome_anexo = nome_anexo[:-4] + "_2.pdf"
                    nomes.append(nome_anexo)
                    anexos.append((nome_anexo, dados))
                assunto, html, texto = self._corpo_email(cfg, cli, comp, protocolo, itens_txt)
                msg = email_srv.montar(ec, para, assunto, html, texto, anexos)
                try:
                    ok, erro = email_srv.enviar(ec, [msg])[0]
                except Exception as e:  # noqa: BLE001
                    ok, erro = False, str(e)
                if not ok:
                    resultados.append({"cliente": cli["nome"], "ok": False, "erro": erro})
                    break  # para no primeiro erro
                hoje = guias.hoje_iso()
                envio = {"protocolo": protocolo, "itens": [i["obrigId"] for i in g["itens"]], "para": para,
                         "assunto": assunto, "anexos": nomes, "enviadoEm": agora_iso(), "threadId": None,
                         "messageId": msg["Message-ID"], "status": "aguardando", "cobrancas": []}
                t = self.banco.obter("tarefas", tid)
                envios = (t.get("envios") or []) + [envio]
                upd_itens = {}
                for i in g["itens"]:
                    atual = (t.get("itens") or {}).get(i["obrigId"])
                    if atual is None:
                        atual = {"venc": i.get("vencimento"), "status": "pendente", "resp": cli.get("resp") or "", "etapas": []}
                    etapas = atual.get("etapas") or []
                    for e in etapas:
                        if e.get("n", "").lower().startswith("enviar ao cliente") and not e.get("ok"):
                            e["ok"], e["em"] = True, hoje
                    atual.update({"status": "enviada", "valor": i.get("valor"), "arquivo": i["arquivo"], "etapas": etapas})
                    atual.setdefault("concluidaEm", hoje)
                    upd_itens[i["obrigId"]] = atual
                self.banco.update("tarefas", tid, {"envios": envios, "itens": upd_itens})
                resultados.append({"cliente": cli["nome"], "ok": True, "protocolo": protocolo, "para": para})
        return {"resultados": resultados}

    # ---------------- ciência e cobrança ----------------
    def conferir_ciencia(self, _corpo):
        ec = self._email_cfg()
        if not ec.get("imap"):
            raise Erro("Informe o servidor IMAP em Configurações > E-mail para ler as respostas.")
        pend = []
        for t in self.banco.listar("tarefas"):
            for e in t.get("envios") or []:
                if e.get("status") != "ciente" and e.get("enviadoEm"):
                    pend.append(e)
        if not pend:
            return {"novos": [], "aguardando": 0}
        achados = email_srv.buscar_respostas(ec, pend)
        novos = []
        for t in self.banco.listar("tarefas"):
            envios = t.get("envios") or []
            mudou, itens = False, {}
            for e in envios:
                r = achados.get(e.get("protocolo"))
                if r and e.get("status") != "ciente":
                    e.update({"status": "ciente", "cienteVia": "Resposta ao e-mail", **r})
                    mudou = True
                    for oid in e.get("itens") or []:
                        if (t.get("itens") or {}).get(oid, {}).get("status") == "enviada":
                            itens[oid] = {"status": "ciente"}
                    alerta = bool(re.search(r"err|errad|não receb|nao receb|reenv|parcel|dúvida|duvida|valor", r["resposta"], re.I))
                    novos.append({"protocolo": e["protocolo"], "cliente": t.get("clienteNome"), "de": r["cienteDe"],
                                  "resposta": r["resposta"], "verificar": alerta})
            if mudou:
                tid = t["id"]
                self.banco.update("tarefas", tid, {"envios": envios, **({"itens": itens} if itens else {})})
        return {"novos": novos, "aguardando": len(pend) - len(novos)}

    def cobrar(self, corpo):
        ec = self._email_cfg()
        cfg = self.banco.obter("config", "escritorio") or {}
        t = self.banco.obter("tarefas", corpo["tarefaId"])
        if not t:
            raise Erro("Tarefa não encontrada.")
        envios = t.get("envios") or []
        e = envios[int(corpo["idx"])]
        obrig = {o["id"]: o for o in self.banco.listar("obrigacoes")}
        linhas = []
        for oid in e.get("itens") or []:
            it = (t.get("itens") or {}).get(oid, {})
            linhas.append(f"{obrig.get(oid, {}).get('sigla', oid)} – vence {guias.br(it.get('venc'))}")
        texto = (f"Olá, {t.get('clienteNome', '')}.\n\nEnviamos em {guias.br(e['enviadoEm'][:10])} as guias abaixo "
                 f"(protocolo {e['protocolo']}) e ainda não recebemos sua confirmação:\n\n"
                 + "\n".join("- " + l for l in linhas)
                 + "\n\nPode responder a este e-mail confirmando o recebimento?\n\n" + (cfg.get("assinatura") or "").replace("{escritorio}", cfg.get("nome") or ""))
        html = "<p>" + texto.replace("&", "&amp;").replace("<", "&lt;").replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>"
        assunto = e.get("assunto") or f"Protocolo {e['protocolo']}"
        if not assunto.lower().startswith("re:"):
            assunto = "Re: " + assunto
        msg = email_srv.montar(ec, e.get("para") or [], assunto, html, texto, responder_a=e.get("messageId"))
        ok, erro = email_srv.enviar(ec, [msg])[0]
        if not ok:
            raise Erro(f"Falha no envio: {erro}")
        e.setdefault("cobrancas", []).append(agora_iso())
        self.banco.update("tarefas", corpo["tarefaId"], {"envios": envios})
        return {"ok": True}

    # ---------------- importação de clientes ----------------
    @staticmethod
    def _linhas_planilha(nome, dados):
        if nome.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            return [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        txt = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                txt = dados.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        dial = csv.Sniffer().sniff(txt[:2000], delimiters=";,\t") if txt else csv.excel
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(txt), dial)]

    @staticmethod
    def _regime(v):
        t = (v or "").lower()
        if "mei" in t or "microempreendedor" in t:
            return "MEI"
        if "simples" in t:
            return "Simples Nacional"
        if "presumido" in t:
            return "Lucro Presumido"
        if "real" in t:
            return "Lucro Real"
        if "imune" in t or "isent" in t:
            return "Imune/Isenta"
        if "física" in t or "fisica" in t or t.strip() == "pf":
            return "Pessoa Física"
        return ""

    def analisar_clientes(self, corpo):
        linhas = self._linhas_planilha(corpo["nome"], base64.b64decode(corpo["b64"]))
        linhas = [l for l in linhas if any(l)]
        if not linhas:
            raise Erro("Planilha vazia.")
        cab = [c.lower() for c in linhas[0]]

        def col(*chaves):
            for i, c in enumerate(cab):
                if any(k in c for k in chaves):
                    return i
            return None

        ic = {"nome": col("razão", "razao", "nome", "empresa", "cliente"), "cnpj": col("cnpj", "cpf", "documento"),
              "regime": col("regime", "tributa"), "email": col("e-mail", "email"), "resp": col("respons"),
              "local": col("munic", "cidade", "uf")}
        if ic["nome"] is None or ic["cnpj"] is None:
            raise Erro("Não achei as colunas de nome e CNPJ na primeira linha da planilha.")
        obrig = self.banco.listar("obrigacoes")
        existentes = {c.get("cnpj"): c for c in self.banco.listar("clientes")}
        vistos, out = set(), []
        for l in linhas[1:]:
            g = lambda k: (l[ic[k]] if ic[k] is not None and ic[k] < len(l) else "").strip()  # noqa: E731
            doc = re.sub(r"\D", "", g("cnpj"))
            if not g("nome") and not doc:
                continue
            if len(doc) in (13, 10):
                doc = doc.zfill(14 if len(doc) == 13 else 11)
            reg = self._regime(g("regime"))
            emails = [e for e in re.split(r"[,;\s]+", g("email")) if "@" in e]
            out.append({"nome": g("nome"), "cnpj": doc, "regime": reg, "emails": emails, "resp": g("resp"),
                        "local": g("local"), "obrigacoes": [o["id"] for o in obrig if reg in (o.get("regimes") or [])],
                        "existe": doc in existentes, "repetido": doc in vistos,
                        "invalido": len(doc) not in (11, 14)})
            vistos.add(doc)
        return {"clientes": out}

    def importar_clientes(self, corpo):
        n = 0
        for c in corpo.get("clientes") or []:
            if c.get("invalido") or c.get("repetido"):
                continue
            id_ = c["cnpj"]
            dados = {k: c[k] for k in ("nome", "cnpj", "regime", "emails", "resp", "local", "obrigacoes")}
            atual = self.banco.obter("clientes", id_)
            if atual:
                # não apaga ajustes feitos na tela: só completa o que está vazio
                for k, v in dados.items():
                    if not atual.get(k):
                        atual[k] = v
                self.banco.set("clientes", id_, atual)
            else:
                dados.update({"ativo": True, "obs": ""})
                self.banco.set("clientes", id_, dados)
            n += 1
        return {"gravados": n}


class Handler(BaseHTTPRequestHandler):
    app: App = None
    protocol_version = "HTTP/1.1"

    def log_message(self, *a):  # silencioso
        pass

    def _json(self, obj, status=200):
        b = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(b)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(b)

    def _corpo(self):
        n = int(self.headers.get("Content-Length") or 0)
        return json.loads(self.rfile.read(n) or b"{}") if n else {}

    def _arquivo(self, caminho):
        if not os.path.isfile(caminho):
            self.send_error(404)
            return
        with open(caminho, "rb") as f:
            b = f.read()
        self.send_response(200)
        tipo = mimetypes.guess_type(caminho)[0] or "application/octet-stream"
        if tipo.startswith("text/") or tipo.endswith("javascript"):
            tipo += "; charset=utf-8"
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(b)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(b)

    def _local(self):
        return self.client_address[0] in ("127.0.0.1", "::1")

    def do_GET(self):
        if not self._local():
            return self.send_error(403)
        u = urlparse(self.path)
        p = unquote(u.path)
        a = self.app
        try:
            if p in ("/", "/index.html"):
                return self._arquivo(os.path.join(WEB, "index.html"))
            if p.startswith("/web/"):
                alvo = os.path.normpath(os.path.join(WEB, p[5:]))
                if not alvo.startswith(WEB):
                    return self.send_error(403)
                return self._arquivo(alvo)
            if p == "/api/rev":
                return self._json({"rev": a.banco.rev()})
            if p == "/api/alteradas":
                desde = int((u.query.split("desde=")[-1] if "desde=" in u.query else "0") or 0)
                return self._json({"rev": a.banco.rev(), "colecoes": a.banco.colecoes_alteradas(desde)})
            m = re.fullmatch(r"/api/col/([a-z]+)", p)
            if m and m.group(1) in COLECOES:
                return self._json(a.banco.listar(m.group(1)))
            m = re.fullmatch(r"/api/doc/([a-z]+)/([^/]+)", p)
            if m and m.group(1) in COLECOES:
                return self._json({"existe": a.banco.obter(*m.groups()) is not None, "data": a.banco.obter(*m.groups())})
            if p == "/api/novoid":
                return self._json({"id": Banco.novo_id()})
            if p == "/api/info":
                return self._json({"versao": a.versao, "pastaDados": a.pasta_dados, "banco": a.banco.caminho})
            if p == "/api/email/config":
                c = a.cfg_email.ler()
                return self._json({"config": c, "provedores": email_srv.PROVEDORES})
            m = re.fullmatch(r"/api/guia/(\d{4}-\d{2})/(.+\.pdf)", p, re.I)
            if m:
                return self._arquivo(os.path.join(a._pasta_guias(m.group(1)), os.path.basename(m.group(2))))
            self.send_error(404)
        except Erro as e:
            self._json({"erro": str(e)}, e.status)
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            self._json({"erro": f"Erro interno: {e}"}, 500)

    def do_POST(self):
        if not self._local():
            return self.send_error(403)
        p = unquote(urlparse(self.path).path)
        a = self.app
        try:
            corpo = self._corpo()
            m = re.fullmatch(r"/api/doc/([a-z]+)/([^/]+)", p)
            if m and m.group(1) in COLECOES:
                col, id_ = m.groups()
                op = corpo.get("op")
                if op == "set":
                    a.banco.set(col, id_, corpo.get("data") or {})
                elif op == "update":
                    try:
                        a.banco.update(col, id_, corpo.get("data") or {})
                    except KeyError:
                        raise Erro("Documento não existe.", 404)
                elif op == "delete":
                    a.banco.delete(col, id_)
                else:
                    raise Erro("Operação inválida.")
                return self._json({"ok": True, "rev": a.banco.rev()})
            rotas = {
                "/api/guias/analisar": a.analisar_guias,
                "/api/guias/enviar": a.enviar_guias,
                "/api/ciencia/conferir": a.conferir_ciencia,
                "/api/cobrar": a.cobrar,
                "/api/clientes/analisar": a.analisar_clientes,
                "/api/clientes/importar": a.importar_clientes,
            }
            if p in rotas:
                return self._json(rotas[p](corpo))
            if p == "/api/email/config":
                a.cfg_email.gravar(corpo)
                return self._json({"ok": True})
            if p == "/api/email/testar":
                c = a.cfg_email.ler(com_senha=True)
                if not c.get("senha"):
                    raise Erro("Salve a senha antes de testar.")
                erros = email_srv.testar(c)
                return self._json({"ok": not erros, "erros": erros})
            if p == "/api/pasta-dados":
                pasta = (corpo.get("pasta") or "").strip()
                if not pasta:
                    raise Erro("Informe a pasta.")
                os.makedirs(pasta, exist_ok=True)
                teste = os.path.join(pasta, ".teste-gravacao")
                with open(teste, "w") as f:
                    f.write("ok")
                os.remove(teste)
                if a.salvar_pasta_dados:
                    a.salvar_pasta_dados(pasta)
                return self._json({"ok": True, "reiniciar": True})
            self.send_error(404)
        except Erro as e:
            self._json({"erro": str(e)}, e.status)
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            self._json({"erro": f"Erro: {e}"}, 500)


def iniciar(app: App, porta=0):
    Handler.app = app
    srv = ThreadingHTTPServer(("127.0.0.1", porta), Handler)
    srv.daemon_threads = True
    th = threading.Thread(target=srv.serve_forever, daemon=True)
    th.start()
    return srv, srv.server_address[1]
